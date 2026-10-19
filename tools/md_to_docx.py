# -*- coding: utf-8 -*-
"""把 Markdown 转成交付用 Word（按本项目既有排版规格）。

为什么自己写而不用 pandoc：环境里没有 pandoc，而且交付要求固定的
中文字体与字号（微软雅黑 / Title 18pt 深蓝 #17375E / 一级标题 15pt /
正文 11pt / A4 / 页边距 72 与 89.85pt）。自己写反而更可控。

只实现交付档需要的子集：标题、段落、无序/有序列表、表格、代码块、
引用、加粗、行内代码。够用即可，不做完整 Markdown。

用法：
    python tools/md_to_docx.py 输入.md 输出.docx "文档标题"
"""

from __future__ import annotations

import io
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "微软雅黑"
MONO = "Consolas"
TITLE_COLOR = RGBColor(0x17, 0x37, 0x5E)


def _set_font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    # 中文必须显式设 eastAsia，否则 Word 会回落成宋体
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(paragraph, hex_fill="F2F2F2"):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


# 彩色 emoji 在微软雅黑里没有字形，会渲染成方框（豆腐块）。
# 交付档统一去掉，语义已在文字里，不缺图标。
_EMOJI = {"✅": "", "❌": "", "📦": "", "🔍": "", "📌": "", "⭐": "",
          "⚠️": "⚠", "❗": "!", "➡️": "→"}


def _clean(text):
    for k, v in _EMOJI.items():
        text = text.replace(k, v)
    return text


def _inline(paragraph, text, size=11, base_bold=False):
    """处理 **加粗** 与 `行内代码`。"""
    text = _clean(text)
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            # 加粗段里可能还套着 `行内代码`，反引号要吃掉，
            # 否则 Word 里会露出 `sym_label` 这样的原样反引号
            r = paragraph.add_run(part[2:-2].replace("`", ""))
            _set_font(r, FONT, size, True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            _set_font(r, MONO, size - 0.5, base_bold)
        else:
            r = paragraph.add_run(part)
            _set_font(r, FONT, size, base_bold)


def _setup(doc, title):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = Pt(72)
    sec.bottom_margin = Pt(72)
    sec.left_margin = Pt(89.85)
    sec.right_margin = Pt(89.85)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    _set_font(r, FONT, 18, True, TITLE_COLOR)
    p.paragraph_format.space_after = Pt(18)


def _table(doc, rows):
    ncol = max(len(r) for r in rows)

    # 列宽按内容长度分配，不是平均分。
    # 平均分的话「名称」列会占掉一大片而「能拿来干什么」窄到每行两个字，
    # 整张表被拉成两页还读不通 —— 这是第一版实拍出来才发现的。
    usable_cm = 14.66
    weights = []
    for j in range(ncol):
        longest = max(len(r[j]) if j < len(r) else 0 for r in rows)
        weights.append(max(4.0, min(float(longest), 42.0)))
    total_w = sum(weights)
    widths = [usable_cm * w / total_w for w in weights]

    # 每列给一个下限，否则短字段（如「许可」）会被压成竖排单字。
    # 补足的部分从最宽的列里扣，保证总宽不变。
    min_cm = 2.0
    for _ in range(8):
        short = [j for j, w in enumerate(widths) if w < min_cm - 1e-6]
        if not short:
            break
        need = sum(min_cm - widths[j] for j in short)
        widest = max(range(ncol), key=lambda j: widths[j])
        widths[widest] -= need
        for j in short:
            widths[j] = min_cm

    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    # 固定布局，否则 Word 会按自己的算法重排，前面算的宽度白算
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    fs = 8.5 if ncol >= 5 else 9
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        # 一行不许跨页拆断：拆开后第二页只剩半个单元格，看不出是什么
        trPr = t.rows[-1]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        if i == 0:
            # 表头跨页重复
            th = OxmlElement("w:tblHeader")
            trPr.append(th)
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            cell = cells[j]
            cell.width = Cm(widths[j])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if i == 0:
                # 表头与后一段同页：否则表格正好卡在页底时，
                # 会出现"只有表头、没有数据行"的孤行（实拍第 1 页就是）
                p.paragraph_format.keep_with_next = True
            _inline(p, txt, size=fs, base_bold=(i == 0))
    doc.add_paragraph()


def convert(md_path, out_path, title):
    with io.open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    _setup(doc, title)

    i = 0
    n = len(lines)
    first_content = True
    while i < n:
        line = lines[i].rstrip()

        # 文首的一级标题就是文档标题，已经单独排过版了，跳过避免重复。
        # （第一版就是这样：封面标题下面又印了一遍一模一样的标题）
        if first_content and not line.strip():
            i += 1
            continue
        if first_content and re.match(r"^#\s+", line):
            first_content = False
            i += 1
            continue
        if line.strip():
            first_content = False

        # 代码块
        if line.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            _shade(p)
            r = p.add_run("\n".join(buf))
            _set_font(r, MONO, 9)
            continue

        # 表格：当前行以 | 开头且下一行是分隔行
        if line.startswith("|") and i + 1 < n and re.match(
                r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                rows.append([c.strip() for c in raw.split("|")])
                i += 1
            rows = [r for k, r in enumerate(rows) if k != 1]  # 去掉 |---|
            _table(doc, rows)
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            p = doc.add_paragraph()
            if level == 1:
                _inline(p, text, size=15, base_bold=True)
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
            elif level == 2:
                _inline(p, text, size=13, base_bold=True)
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
            else:
                _inline(p, text, size=11, base_bold=True)
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # 分隔线
        if re.match(r"^---+$", line):
            i += 1
            continue

        # 引用
        if line.startswith(">"):
            text = line.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _inline(p, text, size=10)
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _inline(p, m.group(2), size=11)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            _inline(p, m.group(3), size=11)
            p.paragraph_format.space_after = Pt(2)
            i += 1
            continue

        # 空行
        if not line.strip():
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        _inline(p, line, size=11)
        i += 1

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("用法：python tools/md_to_docx.py 输入.md 输出.docx 标题")
        raise SystemExit(2)
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                  errors="replace")
    out = convert(sys.argv[1], sys.argv[2], sys.argv[3])
    print("已生成：%s（%.1f KB）" % (out, os.path.getsize(out) / 1024.0))
